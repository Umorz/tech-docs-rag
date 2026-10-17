"""
Microsoft Word (.docx) extractor via python-docx.

DOCX has no hard page numbers (those are computed by the rendering engine),
so we return the entire document as a single page. Paragraphs are joined by
newlines so that the downstream «split into lines» step yields meaningful
line numbers for citations.
"""
from __future__ import annotations

from pathlib import Path

from .. import text_clean
from .base import ExtractedPage, Extractor


class DocxExtractor(Extractor):
    extensions = (".docx",)

    def extract(self, file_path: Path) -> list[ExtractedPage]:
        try:
            from docx import Document  # type: ignore
        except Exception as e:
            raise RuntimeError(
                "python-docx is required for DOCX extraction. "
                "Install with: pip install python-docx. "
                f"Original error: {e}"
            ) from e

        doc = Document(str(file_path))

        parts: list[str] = []
        for paragraph in doc.paragraphs:
            txt = (paragraph.text or "").strip()
            if txt:
                parts.append(txt)

        # Also extract text from tables (cells joined by tabs).
        for table in doc.tables:
            for row in table.rows:
                cells = [(cell.text or "").strip() for cell in row.cells]
                line = "\t".join(c for c in cells if c)
                if line:
                    parts.append(line)

        text = "\n".join(parts)
        text = text_clean.clean_page_text(text)
        return [ExtractedPage(page_no=1, text=text, is_ocr=False)]
